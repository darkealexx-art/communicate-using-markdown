"""Word report generation."""

from __future__ import annotations

from datetime import datetime, timezone
from typing import Iterable

from docx import Document

from news_analyst.analysis import SegmentAnalysis
from news_analyst.config import ReportConfig
from news_analyst.utils import ensure_parent_dir


def _format_date(value: datetime | None) -> str:
    if not value:
        return "No disponible"
    return value.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def _add_list(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_report(report: ReportConfig, analyses: Iterable[SegmentAnalysis]) -> str:
    doc = Document()
    doc.add_heading(report.title, level=0)
    doc.add_paragraph(
        f"Idioma: {report.language} | Periodo: {report.period} | "
        f"Rango: últimos {report.date_range_days} días"
    )
    doc.add_paragraph(f"Generado: {_format_date(datetime.now(timezone.utc))}")

    for analysis in analyses:
        doc.add_heading(analysis.segment_display_name, level=1)
        doc.add_heading("Resumen ejecutivo", level=2)
        doc.add_paragraph(analysis.executive_summary)

        doc.add_heading("Análisis profundo", level=2)
        doc.add_paragraph(analysis.deep_analysis)

        doc.add_heading("Tendencias", level=2)
        if analysis.trends:
            _add_list(doc, analysis.trends)
        else:
            doc.add_paragraph("Sin tendencias destacadas.")

        doc.add_heading("Riesgos", level=2)
        if analysis.risks:
            _add_list(doc, analysis.risks)
        else:
            doc.add_paragraph("No se identificaron riesgos críticos.")

        doc.add_heading("Oportunidades", level=2)
        if analysis.opportunities:
            _add_list(doc, analysis.opportunities)
        else:
            doc.add_paragraph("No se identificaron oportunidades relevantes.")

        doc.add_heading("Implicaciones", level=2)
        _add_list(doc, analysis.implications)

        doc.add_heading("Noticias destacadas", level=2)
        if not analysis.items:
            doc.add_paragraph("Sin noticias disponibles.")
        for item in analysis.items:
            doc.add_paragraph(
                f"{item.title} ({item.source}) - {item.impact_level.upper()} "
                f"Relevancia {item.relevance_score}/5",
                style="List Bullet",
            )
            doc.add_paragraph(
                f"Fecha: {_format_date(item.published_at)} | URL: {item.url}"
            )
            if item.summary:
                doc.add_paragraph(f"Resumen: {item.summary}")

        doc.add_heading("Fuentes consultadas", level=2)
        sources = {
            f"{item.source} - {item.url}" for item in analysis.items
        }
        if sources:
            _add_list(doc, sorted(sources))
        else:
            doc.add_paragraph("Sin fuentes disponibles.")

    output_path = ensure_parent_dir(report.output_path)
    doc.save(output_path)
    return str(output_path)
